import json
import re
from pathlib import Path
from typing import Any

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import markdown
except ImportError:
    markdown = None

try:
    from docx import Document
except ImportError:
    Document = None


def parse_pdf(file_path: Path) -> dict[str, Any]:
    """解析PDF文件"""
    if fitz is None:
        raise ImportError("PyMuPDF not installed. Run: pip install PyMuPDF")

    doc = fitz.open(file_path)
    total_pages = len(doc)
    chapters = []
    current_chapter = None
    chapter_counter = 0

    for page_num in range(total_pages):
        page = doc[page_num]
        text = page.get_text()

        # 简单的章节识别：查找"第X章"或"Chapter X"
        chapter_match = re.search(r'第[一二三四五六七八九十\d]+章\s*(.+)', text)
        if not chapter_match:
            chapter_match = re.search(r'Chapter\s+(\d+)[:\s]+(.+)', text, re.IGNORECASE)

        if chapter_match:
            # 保存上一章
            if current_chapter:
                chapters.append(current_chapter)

            chapter_counter += 1
            chapter_title = chapter_match.group(0).strip()
            current_chapter = {
                "chapter_id": f"ch_{chapter_counter:02d}",
                "title": chapter_title,
                "page_start": page_num + 1,
                "page_end": page_num + 1,
                "content": text,
                "char_count": len(text),
            }
        elif current_chapter:
            # 继续当前章节
            current_chapter["content"] += "\n" + text
            current_chapter["page_end"] = page_num + 1
            current_chapter["char_count"] += len(text)

    # 保存最后一章
    if current_chapter:
        chapters.append(current_chapter)

    # 如果没有识别到章节，将整个文档作为一章
    if not chapters:
        all_text = ""
        for page_num in range(total_pages):
            all_text += doc[page_num].get_text() + "\n"

        chapters.append({
            "chapter_id": "ch_01",
            "title": "全文",
            "page_start": 1,
            "page_end": total_pages,
            "content": all_text,
            "char_count": len(all_text),
        })

    doc.close()

    return {
        "total_pages": total_pages,
        "total_chars": sum(ch["char_count"] for ch in chapters),
        "chapters": chapters,
    }


def parse_markdown(file_path: Path) -> dict[str, Any]:
    """解析Markdown文件"""
    content = file_path.read_text(encoding='utf-8')

    # 按标题分章节
    chapters = []
    lines = content.split('\n')
    current_chapter = None
    chapter_counter = 0

    for line in lines:
        # 识别一级或二级标题
        if line.startswith('# ') or line.startswith('## '):
            if current_chapter:
                chapters.append(current_chapter)

            chapter_counter += 1
            title = line.lstrip('#').strip()
            current_chapter = {
                "chapter_id": f"ch_{chapter_counter:02d}",
                "title": title,
                "page_start": chapter_counter,
                "page_end": chapter_counter,
                "content": line + "\n",
                "char_count": len(line),
            }
        elif current_chapter:
            current_chapter["content"] += line + "\n"
            current_chapter["char_count"] += len(line)

    if current_chapter:
        chapters.append(current_chapter)

    if not chapters:
        chapters.append({
            "chapter_id": "ch_01",
            "title": "全文",
            "page_start": 1,
            "page_end": 1,
            "content": content,
            "char_count": len(content),
        })

    return {
        "total_pages": len(chapters),
        "total_chars": len(content),
        "chapters": chapters,
    }


def parse_txt(file_path: Path) -> dict[str, Any]:
    """解析TXT文件"""
    content = file_path.read_text(encoding='utf-8')

    # 简单按段落分章节
    chapters = []
    paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]

    for i, para in enumerate(paragraphs[:10], 1):  # 最多取前10段作为章节
        chapters.append({
            "chapter_id": f"ch_{i:02d}",
            "title": f"第{i}段",
            "page_start": i,
            "page_end": i,
            "content": para,
            "char_count": len(para),
        })

    if not chapters:
        chapters.append({
            "chapter_id": "ch_01",
            "title": "全文",
            "page_start": 1,
            "page_end": 1,
            "content": content,
            "char_count": len(content),
        })

    return {
        "total_pages": 1,
        "total_chars": len(content),
        "chapters": chapters,
    }


def parse_docx(file_path: Path) -> dict[str, Any]:
    """解析Word文档"""
    if Document is None:
        raise ImportError("python-docx not installed. Run: pip install python-docx")

    doc = Document(file_path)
    chapters = []
    current_chapter = None
    chapter_counter = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # 识别标题样式或章节标记
        if para.style.name.startswith('Heading') or re.search(r'第[一二三四五六七八九十\d]+章', text):
            if current_chapter:
                chapters.append(current_chapter)

            chapter_counter += 1
            current_chapter = {
                "chapter_id": f"ch_{chapter_counter:02d}",
                "title": text,
                "page_start": chapter_counter,
                "page_end": chapter_counter,
                "content": text + "\n",
                "char_count": len(text),
            }
        elif current_chapter:
            current_chapter["content"] += text + "\n"
            current_chapter["char_count"] += len(text)

    if current_chapter:
        chapters.append(current_chapter)

    if not chapters:
        all_text = "\n".join([p.text for p in doc.paragraphs])
        chapters.append({
            "chapter_id": "ch_01",
            "title": "全文",
            "page_start": 1,
            "page_end": 1,
            "content": all_text,
            "char_count": len(all_text),
        })

    return {
        "total_pages": len(chapters),
        "total_chars": sum(ch["char_count"] for ch in chapters),
        "chapters": chapters,
    }


def parse_textbook(payload: dict[str, Any]) -> dict[str, Any]:
    """解析教材文件"""
    textbook_id = payload.get("textbook_id", "demo-textbook")
    filename = payload.get("filename", "")
    file_path_str = payload.get("file_path")

    if not file_path_str:
        return {
            "success": False,
            "textbook_id": textbook_id,
            "filename": filename,
            "error": "Missing file_path",
        }

    file_path = Path(file_path_str)
    if not file_path.exists():
        return {
            "success": False,
            "textbook_id": textbook_id,
            "filename": filename,
            "error": f"File not found: {file_path}",
        }

    try:
        # 根据文件扩展名选择解析器
        ext = file_path.suffix.lower()

        if ext == '.pdf':
            parse_result = parse_pdf(file_path)
        elif ext in ['.md', '.markdown']:
            parse_result = parse_markdown(file_path)
        elif ext == '.txt':
            parse_result = parse_txt(file_path)
        elif ext == '.docx':
            parse_result = parse_docx(file_path)
        else:
            return {
                "success": False,
                "textbook_id": textbook_id,
                "filename": filename,
                "error": f"Unsupported file format: {ext}",
            }

        # 保存解析结果
        result = {
            "success": True,
            "textbook_id": textbook_id,
            "filename": filename,
            "title": filename,
            "format": ext.lstrip('.'),
            "total_pages": parse_result["total_pages"],
            "total_chars": parse_result["total_chars"],
            "chapters": parse_result["chapters"],
        }

        # 保存到JSON文件
        data_dir = file_path.parent.parent / "parsed"
        data_dir.mkdir(exist_ok=True)
        output_file = data_dir / f"{textbook_id}.json"
        output_file.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')

        return result

    except Exception as e:
        return {
            "success": False,
            "textbook_id": textbook_id,
            "filename": filename,
            "error": str(e),
        }
